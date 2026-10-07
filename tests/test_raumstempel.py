# -*- coding: utf-8 -*-
"""Tests der Revit-freien Logik. Start: python3 -m unittest discover -s tests -v"""
import math
import os
import sys
import tempfile
import unittest

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HIER, "..", "raumstempel"))
import raumstempel_dynamo as rs  # noqa: E402

BEISPIEL_CSV = os.path.join(HIER, "daten", "100049_004_A_G03_Test.csv")


class TestZahlen(unittest.TestCase):
    def test_punkt_und_komma(self):
        self.assertAlmostEqual(rs.parse_zahl("12.9694"), 12.9694)
        self.assertAlmostEqual(rs.parse_zahl("12,9694"), 12.9694)
        self.assertAlmostEqual(rs.parse_zahl("1.234,56"), 1234.56)
        self.assertAlmostEqual(rs.parse_zahl("1,234.56"), 1234.56)
        self.assertAlmostEqual(rs.parse_zahl(" 4.01 "), 4.01)

    def test_ungueltig(self):
        self.assertIsNone(rs.parse_zahl(""))
        self.assertIsNone(rs.parse_zahl("abc"))
        self.assertIsNone(rs.parse_zahl(None))


class TestCsv(unittest.TestCase):
    def test_beispieldatei(self):
        stempel, mel = rs.lese_stempel_ordner(os.path.dirname(BEISPIEL_CSV))
        self.assertEqual(mel, [])
        self.assertEqual(len(stempel), 23)
        s = [x for x in stempel if x.oks.endswith("G03-_04")][0]
        self.assertEqual(s.name, "Aufzug")
        self.assertEqual(s.nummer, "n.v.")
        self.assertAlmostEqual(s.x, 17.125)
        self.assertAlmostEqual(s.y, 8.9281)
        self.assertAlmostEqual(s.flaeche, 3.29)
        self.assertEqual(s.dateiname, "100049_004_A_G03_Test.dwg")
        self.assertEqual({x.code for x in stempel}, {"G03"})

    def test_umlaute_und_nummer_ohne_leerzeichen(self):
        stempel, _ = rs.lese_stempel_ordner(os.path.dirname(BEISPIEL_CSV))
        self.assertIn("Teeküche", [s.name for s in stempel])
        b = [s for s in stempel if s.oks.endswith("G03-_01")][0]
        self.assertEqual(b.nummer, "4.01")           # Leerzeichen am Ende entfernt

    def test_semikolon_dezimalkomma_cp1252(self):
        text = "FM.OKS;FM.NUMMER;FM.NAME;FM.FLAECHE;Position X;Position Y\r\n" \
               "100049-004-A-G01-_15;2.10b;Büro;17,89;12,5000;3,25\r\n"
        st, mel = rs.parse_stempel_text(rs.dekodiere(text.encode("cp1252")), "t.csv")
        self.assertEqual(mel, [])
        self.assertEqual(st[0].name, "Büro")
        self.assertAlmostEqual(st[0].x, 12.5)
        self.assertAlmostEqual(st[0].flaeche, 17.89)

    def test_fehlende_spalte_und_leerzeile_und_fehlerzeile(self):
        st, mel = rs.parse_stempel_text("a;b\n1;2\n", "x.csv")
        self.assertEqual(st, [])
        self.assertIn("Pflichtspalte", mel[0])
        text = "FM.OKS,Position X,Position Y\n,1,2\nA-G01-_1,x,2\n\nA-G01-_2,1,2\n"
        st, mel = rs.parse_stempel_text(text, "y.csv")
        self.assertEqual([s.oks for s in st], ["A-G01-_2"])
        self.assertEqual(len(mel), 2)

    def test_ziel_koordinaten(self):
        text = "FM.OKS,Position X,Position Y,Ziel X,Ziel Y\nA-G01-_1,1,2,5,6\n"
        st, _ = rs.parse_stempel_text(text, "z.csv")
        self.assertEqual((st[0].punkt_x, st[0].punkt_y), (5.0, 6.0))


class TestExterneReferenzen(unittest.TestCase):
    ORDNER = os.path.join(HIER, "daten", "alle")

    def test_filter_funktion(self):
        self.assertTrue(rs.ist_externe_referenz("100049_004_A_G03_Bestand.dwg"))
        self.assertTrue(rs.ist_externe_referenz("x_BESTAND"))
        self.assertFalse(rs.ist_externe_referenz("100049_004_A_G03.dwg"))
        self.assertFalse(rs.ist_externe_referenz("Bestand_G03.dwg"))
        self.assertFalse(rs.ist_externe_referenz("", "_Bestand"))
        self.assertFalse(rs.ist_externe_referenz("a_Bestand.dwg", ""))

    def test_alle_geschosse_datei(self):
        stempel, mel = rs.lese_stempel_ordner(self.ORDNER)
        self.assertEqual(len(stempel), 156)                  # 214 Zeilen - 58 aus Referenzen
        self.assertTrue(all("_Bestand" not in s.dateiname for s in stempel))
        self.assertEqual(len(mel), 1)
        self.assertIn("58 Zeilen aus externen Referenzen ignoriert", mel[0])
        je_code = {}
        for s in stempel:
            je_code[s.code] = je_code.get(s.code, 0) + 1
        self.assertEqual(je_code, {"G00": 33, "G01": 31, "G02": 34, "G03": 23, "G04": 2, "U01": 33})

    def test_danach_alles_eindeutig(self):
        stempel, _ = rs.lese_stempel_ordner(self.ORDNER)
        self.assertEqual(rs.finde_doppelte({i: s.oks for i, s in enumerate(stempel)}), {})
        nummern = {s.oks: rs.kurz_nummer(s.oks)[0] for s in stempel}
        self.assertNotIn(None, nummern.values())
        self.assertEqual(rs.finde_doppelte(nummern), {})

    def test_ohne_filter_gibt_es_doppelte_oks(self):
        stempel, mel = rs.lese_stempel_ordner(self.ORDNER, ausschluss_suffix="")
        self.assertEqual(len(stempel), 214)
        self.assertEqual(mel, [])
        self.assertGreater(len(rs.finde_doppelte({i: s.oks for i, s in enumerate(stempel)})), 0)

    def test_ebenenpruefung_mit_realen_codes(self):
        stempel, _ = rs.lese_stempel_ordner(self.ORDNER)
        codes = {}
        for s in stempel:
            codes[s.code] = codes.get(s.code, 0) + 1
        zuord = rs.parse_ebenen_zuordnung(rs.STANDARD_EBENEN)
        # 1. UG, EG, 1. OG, 2. OG, 3. OG modelliert; 4. OG nicht
        erg = {e["code"]: e for e in rs.pruefe_ebenen(
            codes, zuord, ["1. UG", "EG", "1. OG", "2. OG", "3. OG", "4. OG"],
            {"1. UG": 5, "EG": 5, "1. OG": 5, "2. OG": 5, "3. OG": 5})}
        self.assertEqual([c for c, e in erg.items() if not e["ok"]], ["G04"])


class TestDoppelte(unittest.TestCase):
    def _s(self, oks, name="Büro", x=1.0, y=2.0, fl=10.0, nr="1.01", quelle="a.csv"):
        return rs.Stempel(oks, nummer=nr, name=name, flaeche=fl, x=x, y=y, quelle=quelle)

    def test_identische_kopien_werden_einmal_verwendet(self):
        a = [self._s("A-G01-_1", quelle="Test.csv"), self._s("A-G01-_2", quelle="Test.csv")]
        b = [self._s("A-G01-_1", quelle="Test_voll.csv"), self._s("A-G01-_2", quelle="Test_voll.csv")]
        liste, entfernt, konflikte = rs.bereinige_doppelte(a + b)
        self.assertEqual([s.oks for s in liste], ["A-G01-_1", "A-G01-_2"])
        self.assertEqual(len(entfernt), 2)
        self.assertEqual(konflikte, {})
        self.assertEqual({s.quelle for s in entfernt}, {"Test_voll.csv"})

    def test_widerspruechliche_werden_nicht_verarbeitet(self):
        liste, entfernt, konflikte = rs.bereinige_doppelte(
            [self._s("A-G01-_1"), self._s("A-G01-_1", name="Flur"), self._s("A-G01-_2")])
        self.assertEqual([s.oks for s in liste], ["A-G01-_2"])
        self.assertEqual(list(konflikte), ["A-G01-_1"])
        self.assertEqual(entfernt, [])

    def test_leerzeichen_und_kleine_positionsabweichung_gelten_als_identisch(self):
        liste, entfernt, _ = rs.bereinige_doppelte(
            [self._s("A-G01-_1", name="Büro "), self._s("A-G01-_1", name="Büro", x=1.0001)])
        self.assertEqual(len(liste), 1)
        self.assertEqual(len(entfernt), 1)

    def test_zwei_csv_mit_gleichem_inhalt_im_ordner(self):
        with tempfile.TemporaryDirectory() as d:
            with open(BEISPIEL_CSV, "rb") as quelle:
                inhalt = quelle.read()
            for name in ("Test.csv", "Test_voll.csv"):
                with open(os.path.join(d, name), "wb") as f:
                    f.write(inhalt)
            stempel, _ = rs.lese_stempel_ordner(d)
            self.assertEqual(len(stempel), 46)
            liste, entfernt, konflikte = rs.bereinige_doppelte(stempel)
            self.assertEqual((len(liste), len(entfernt), konflikte), (23, 23, {}))


class TestVerknuepfung(unittest.TestCase):
    """Stempel aus '<G>.dwg', verknüpft ist in Revit '<G>_Bestand.dwg'."""

    def test_norm(self):
        self.assertEqual(rs.norm_dwgname("U:\\x\\100049_004_A_G03_Bestand.DWG"), "100049_004_a_g03_bestand")
        self.assertEqual(rs.norm_dwgname("a/b/G01.dwg"), "g01")
        self.assertEqual(rs.norm_dwgname(None), "")

    def test_kandidaten(self):
        self.assertEqual(rs.kandidaten_linknamen("100049_004_A_G03.dwg"),
                         ["100049_004_a_g03_bestand", "100049_004_a_g03"])
        self.assertEqual(rs.kandidaten_linknamen("G03.dwg", ""), ["g03"])
        self.assertEqual(rs.kandidaten_linknamen("G03.dwg", manuell={"g03": "mein_link"}), ["mein_link"])

    def test_bestand_link_wird_bevorzugt(self):
        name, hinweis = rs.waehle_verknuepfung(
            ["100049_004_A_G03.dwg"],
            ["100049_004_A_G03_Bestand.dwg", "100049_004_A_G03.dwg", "Anderes.dwg"])
        self.assertEqual(name, "100049_004_A_G03_Bestand.dwg")
        self.assertIn("Koordinaten von Verknüpfung", hinweis)

    def test_nur_bestand_link_vorhanden(self):
        name, _ = rs.waehle_verknuepfung(["100049_004_A_G00.dwg"], ["100049_004_A_G00_Bestand.dwg"])
        self.assertEqual(name, "100049_004_A_G00_Bestand.dwg")

    def test_rueckfall_auf_gleichen_namen(self):
        name, hinweis = rs.waehle_verknuepfung(["G01.dwg"], ["G01.dwg"])
        self.assertEqual((name, hinweis), ("G01.dwg", ""))

    def test_kein_treffer(self):
        self.assertEqual(rs.waehle_verknuepfung(["G05.dwg"], ["G01_Bestand.dwg"]), (None, ""))

    def test_manuelle_zuordnung(self):
        name, _ = rs.waehle_verknuepfung(["G05.dwg"], ["Architektur_5OG.dwg"],
                                         manuell={"g05": "architektur_5og"})
        self.assertEqual(name, "Architektur_5OG.dwg")

    def test_alle_geschosse_der_beispieldatei_finden_ihren_link(self):
        stempel, _ = rs.lese_stempel_ordner(os.path.join(HIER, "daten", "alle"))
        links = ["100049_004_A_%s_Bestand.dwg" % c for c in ("G00", "G01", "G02", "G03", "G04", "U01")]
        for code in ("G00", "G01", "G02", "G03", "G04", "U01"):
            dn = [s.dateiname for s in stempel if s.code == code]
            name, _ = rs.waehle_verknuepfung(dn, links)
            self.assertEqual(name, "100049_004_A_%s_Bestand.dwg" % code)


class TestReferenzzeilenNurBestand(unittest.TestCase):
    def test_geschoss_nur_mit_referenzzeilen(self):
        text = ("FM.OKS,Position X,Position Y,Dateiname\n"
                "A-G05-_1,1,2,A_G05_Bestand.dwg\n"
                "A-G06-_1,1,2,A_G06.dwg\n"
                "A-G06-_2,1,2,A_G06_Bestand.dwg\n")
        ign = {}
        st, mel = rs.parse_stempel_text(text, "x.csv", ignoriert_out=ign)
        self.assertEqual([s.oks for s in st], ["A-G06-_1"])
        self.assertEqual(ign, {"G05": 1, "G06": 1})
        self.assertEqual(len(mel), 1)
        self.assertEqual(set(ign) - {s.code for s in st}, {"G05"})


class TestNummern(unittest.TestCase):
    def test_kurz_segmente(self):
        self.assertEqual(rs.kurz_nummer("100049-004-A-G01-_15"), ("G01-_15", None))
        self.assertEqual(rs.kurz_nummer("100049-004-A-G01 -_15"), ("G01-_15", None))
        self.assertEqual(rs.kurz_nummer("100049-004-A-U01-_105"), ("U01-_105", None))

    def test_kurz_laenge(self):
        self.assertEqual(rs.kurz_nummer("100049-004-A-G01-_15", "laenge", 7), ("G01-_15", None))
        n, f = rs.kurz_nummer("G01_15", "laenge", 7)
        self.assertIsNone(n)
        self.assertIn("kürzer", f)

    def test_kurz_fehler(self):
        self.assertIsNone(rs.kurz_nummer("ABC")[0])
        self.assertIsNone(rs.kurz_nummer("100049-004-A-XX-_15")[0])

    def test_beispieldaten_eindeutig(self):
        stempel, _ = rs.lese_stempel_ordner(os.path.dirname(BEISPIEL_CSV))
        nummern = {s.oks: rs.kurz_nummer(s.oks)[0] for s in stempel}
        self.assertTrue(all(nummern.values()))
        self.assertEqual(rs.finde_doppelte(nummern), {})
        self.assertEqual(sorted(nummern.values())[0], "G03-_01")

    def test_doppelte(self):
        d = rs.finde_doppelte({"a": "G01-_1", "b": "G01-_1", "c": "G01-_2"})
        self.assertEqual(d, {"G01-_1": ["a", "b"]})

    def test_geschosscode(self):
        self.assertEqual(rs.geschosscode("100049-004-A-G03-_01"), "G03")
        self.assertEqual(rs.geschosscode("100049-004-A-u01-_01"), "U01")
        self.assertIsNone(rs.geschosscode("100049-004-A-X03-_01"))
        self.assertIsNone(rs.geschosscode(""))


class TestEbenen(unittest.TestCase):
    ZUORD = rs.parse_ebenen_zuordnung(rs.STANDARD_EBENEN)
    REVIT = ["1. UG", "EG", "1. OG", "2. OG", "3. OG", "4.OG"]

    def test_parse(self):
        self.assertEqual(self.ZUORD["G01"], "1. OG")
        self.assertEqual(self.ZUORD["U01"], "1. UG")

    def test_pruefung(self):
        codes = {"G00": 20, "G03": 23, "G04": 5, "U01": 7, "G09": 2, None: 1}
        raeume = {"EG": 30, "3. OG": 25, "4.OG": 0}      # 1. UG nicht modelliert
        erg = {e["code"]: e for e in rs.pruefe_ebenen(codes, self.ZUORD, self.REVIT, raeume)}
        self.assertTrue(erg["G00"]["ok"])
        self.assertEqual(erg["G03"]["ebene"], "3. OG")
        self.assertFalse(erg["G04"]["ok"])               # "4.OG" ohne Räume
        self.assertIn("nicht modelliert", erg["G04"]["status"])
        self.assertFalse(erg["U01"]["ok"])
        self.assertFalse(erg["G09"]["ok"])
        self.assertIn("nicht in der Ebenenzuordnung", erg["G09"]["status"])
        self.assertFalse(erg[None]["ok"])

    def test_namensvarianten(self):
        self.assertEqual(rs.norm_ebenenname("4.OG"), rs.norm_ebenenname("4. OG"))
        erg = rs.pruefe_ebenen({"G04": 1}, self.ZUORD, ["4.OG"], {"4.OG": 3})
        self.assertTrue(erg[0]["ok"])

    def test_ebene_fehlt_im_modell(self):
        erg = rs.pruefe_ebenen({"G01": 3}, self.ZUORD, ["EG"], {"EG": 10})
        self.assertIn("nicht im Revit-Modell", erg[0]["status"])


class TestKoordinaten(unittest.TestCase):
    IDENT = {"origin": (0, 0, 0), "basis_x": (1, 0, 0), "basis_y": (0, 1, 0), "basis_z": (0, 0, 1)}

    def test_faktoren(self):
        self.assertAlmostEqual(rs.einheit_in_fuss("m"), 1 / 0.3048)
        self.assertAlmostEqual(rs.einheit_in_fuss("cm"), 0.01 / 0.3048)
        self.assertAlmostEqual(rs.einheit_in_fuss("mm"), 0.001 / 0.3048)
        with self.assertRaises(ValueError):
            rs.einheit_in_fuss("km")

    def test_nur_einheit(self):
        x, y, z = rs.dwg_nach_revit(10.0, 5.0, "m", self.IDENT)
        self.assertAlmostEqual(x, 10 / 0.3048)
        self.assertAlmostEqual(y, 5 / 0.3048)
        self.assertAlmostEqual(z, 0.0)
        x, y, _ = rs.dwg_nach_revit(10000.0, 5000.0, "mm", self.IDENT)
        self.assertAlmostEqual(x, 10 / 0.3048)

    def test_versatz(self):
        t = dict(self.IDENT, origin=(100.0, 200.0, 5.0))
        x, y, z = rs.dwg_nach_revit(1.0, 2.0, "m", t)
        self.assertAlmostEqual(x, 100.0 + 1 / 0.3048)
        self.assertAlmostEqual(y, 200.0 + 2 / 0.3048)
        self.assertAlmostEqual(z, 5.0)

    def test_drehung_90_grad(self):
        t = {"origin": (0, 0, 0), "basis_x": (0, 1, 0), "basis_y": (-1, 0, 0), "basis_z": (0, 0, 1)}
        x, y, _ = rs.dwg_nach_revit(1.0, 0.0, "m", t)
        self.assertAlmostEqual(x, 0.0)
        self.assertAlmostEqual(y, 1 / 0.3048)

    def test_pruefpunkt(self):
        self.assertAlmostEqual(rs.pruefpunkt_hoehe_fuss(10.0), 10.0 + 1 / 0.3048)


def _st(name, flaeche, oks="A-G03-_%02d", nr=[0]):
    nr[0] += 1
    return rs.Stempel(oks % nr[0], nummer="n.v.", name=name, flaeche=flaeche, x=0, y=0)


class TestZuordnung(unittest.TestCase):
    def test_treppenhaus_mit_drei_stempeln(self):
        """Der Fall aus dem Projekt: 3 Stempel liegen im Treppenhaus-Raum."""
        stempel = [_st("Treppenhaus 1", 48.37), _st("Aufzug", 3.29), _st("nicht bekannt", 1.17)]
        raeume = {1: {"flaeche": 48.2, "pos": (0, 0)},      # Treppenhaus
                  2: {"flaeche": 3.31, "pos": (4, 0)},      # Aufzug
                  3: {"flaeche": 1.2, "pos": (4, 3)}}       # Schacht
        treffer = {0: [1], 1: [1], 2: [1]}
        pos = {0: (0, 0), 1: (1, 0), 2: (1, 1)}
        e = rs.ordne_zu(stempel, raeume, treffer, pos)
        zu = {z.stempel_idx: z for z in e["zuordnungen"]}
        self.assertEqual(zu[0].raum_id, 1)
        self.assertEqual(zu[0].methode, "Position+Fläche")
        self.assertEqual(zu[1].raum_id, 2)
        self.assertEqual(zu[1].methode, "Fläche")
        self.assertEqual(zu[2].raum_id, 3)
        self.assertTrue(all(z.status == "sicher" for z in zu.values()))
        self.assertEqual(e["stempel_ohne_raum"], [])
        self.assertEqual(e["raeume_ohne_stempel"], [])
        self.assertEqual(e["raeume_mehrfach"], {1: [0, 1, 2]})

    def test_einfach_position(self):
        stempel = [_st("Büro", 17.9)]
        e = rs.ordne_zu(stempel, {7: {"flaeche": 17.5, "pos": None}}, {0: [7]})
        z = e["zuordnungen"][0]
        self.assertEqual((z.raum_id, z.methode, z.status), (7, "Position", "sicher"))

    def test_flaeche_weicht_stark_ab_ist_unsicher(self):
        stempel = [_st("Büro", 17.9)]
        e = rs.ordne_zu(stempel, {7: {"flaeche": 30.0, "pos": None}}, {0: [7]})
        z = e["zuordnungen"][0]
        self.assertEqual(z.status, "unsicher")
        self.assertIn("weicht", z.bemerkung)

    def test_mittlere_abweichung_unsicher(self):
        stempel = [_st("Büro", 17.9)]
        e = rs.ordne_zu(stempel, {7: {"flaeche": 17.9 / 0.90, "pos": None}}, {0: [7]})  # ~11 %
        self.assertEqual(e["zuordnungen"][0].status, "unsicher")

    def test_aehnliche_flaechen_abstand_entscheidet_unsicher(self):
        """Mehrere ähnlich große Büros, Stempel liegen außerhalb: Nähe entscheidet, aber unsicher."""
        stempel = [_st("Büro A", 17.9), _st("Büro B", 17.6)]
        raeume = {1: {"flaeche": 17.8, "pos": (0, 0)}, 2: {"flaeche": 17.7, "pos": (10, 0)}}
        pos = {0: (1, 0), 1: (9, 0)}
        e = rs.ordne_zu(stempel, raeume, {0: [], 1: []}, pos)
        zu = {z.stempel_idx: z for z in e["zuordnungen"]}
        self.assertEqual((zu[0].raum_id, zu[1].raum_id), (1, 2))
        self.assertTrue(all(z.status == "unsicher" for z in zu.values()))

    def test_max_abstand(self):
        stempel = [_st("Büro", 17.9)]
        raeume = {1: {"flaeche": 17.9, "pos": (100, 0)}}
        e = rs.ordne_zu(stempel, raeume, {0: []}, {0: (0, 0)}, max_abstand=10)
        self.assertEqual(e["zuordnungen"], [])
        self.assertEqual(e["stempel_ohne_raum"], [0])
        self.assertEqual(e["raeume_ohne_stempel"], [1])

    def test_stempel_ohne_raum_und_raum_ohne_stempel(self):
        stempel = [_st("X", 40.0)]
        raeume = {1: {"flaeche": 10.0, "pos": (0, 0)}}
        e = rs.ordne_zu(stempel, raeume, {0: []}, {0: (0, 0)})
        self.assertEqual(e["stempel_ohne_raum"], [0])
        self.assertEqual(e["raeume_ohne_stempel"], [1])

    def test_ohne_stempelflaeche(self):
        stempel = [_st("Büro", None)]
        e = rs.ordne_zu(stempel, {7: {"flaeche": 17.5, "pos": None}}, {0: [7]})
        z = e["zuordnungen"][0]
        self.assertEqual(z.status, "unsicher")
        self.assertIsNone(z.abweichung)

    def test_stempel_in_fremdem_raum_wird_per_flaeche_umgeleitet(self):
        stempel = [_st("Teeküche", 2.83)]
        raeume = {1: {"flaeche": 40.0, "pos": (0, 0)}, 2: {"flaeche": 2.8, "pos": (3, 0)}}
        e = rs.ordne_zu(stempel, raeume, {0: [1]}, {0: (1, 0)})
        z = e["zuordnungen"][0]
        self.assertEqual(z.raum_id, 2)
        self.assertEqual(e["raeume_ohne_stempel"], [1])


class TestSchreiben(unittest.TestCase):
    def test_schreibvorgaben(self):
        s = rs.Stempel("100049-004-A-G01-_15", nummer="4.06b ", name=" Büro ", flaeche=1, x=0, y=0)
        w, fehler = rs.schreibvorgaben(s)
        self.assertIsNone(fehler)
        self.assertEqual(w, {"name": "Büro", "oks": "100049-004-A-G01-_15",
                             "nummer_text": "Raum-Nr. 4.06b", "nummer": "G01-_15"})

    def test_raumname_unveraendert_auch_nv(self):
        s = rs.Stempel("100049-004-A-G01-_31", nummer="n.v.", name="n.v.", x=0, y=0)
        self.assertEqual(rs.schreibvorgaben(s)[0]["name"], "n.v.")
        s = rs.Stempel("100049-004-A-G01-_31", nummer="1", name="Aufenthaltsraum ", x=0, y=0)
        self.assertEqual(rs.schreibvorgaben(s)[0]["name"], "Aufenthaltsraum")   # nur Randleerzeichen

    def test_nv_bleibt_erhalten(self):
        s = rs.Stempel("100049-004-A-G01-_15", nummer="n.v.", name="Flur", x=0, y=0)
        self.assertEqual(rs.schreibvorgaben(s)[0]["nummer_text"], "Raum-Nr. n.v.")

    def test_nur_aenderungen(self):
        soll = {"name": "Büro", "oks": "X", "nummer_text": "Raum-Nr. 1", "nummer": "G01-_1"}
        ist = {"name": "Büro", "oks": "alt", "nummer_text": "Raum-Nr. 1", "nummer": None}
        self.assertEqual(rs.aenderungen(ist, soll), {"oks": "X", "nummer": "G01-_1"})
        self.assertEqual(rs.aenderungen(soll, soll), {})          # zweiter Lauf: nichts zu tun

    def test_none_wird_nicht_geschrieben(self):
        self.assertEqual(rs.aenderungen({"nummer": "a"}, {"nummer": None}), {})


class TestEingaben(unittest.TestCase):
    def test_liste_in_in0(self):
        e = rs._eingaben_entpacken([["C:/x", True, False]])
        self.assertEqual(rs._eingabe(e, 0, ""), "C:/x")
        self.assertEqual(rs._eingabe(e, 1, None), True)
        self.assertEqual(rs._eingabe(e, 9, "Raum-Nr. "), "Raum-Nr. ")    # fehlt -> Standard

    def test_einzelne_eingaenge(self):
        e = rs._eingaben_entpacken(["C:/x", True])
        self.assertEqual(rs._eingabe(e, 0, ""), "C:/x")
        self.assertEqual(rs._eingabe(e, 5, "Bestand"), "Bestand")
        self.assertEqual(rs._eingabe([None, ""], 1, "std"), "std")


class TestListen(unittest.TestCase):
    def test_csv_roundtrip_und_freigabe(self):
        with tempfile.TemporaryDirectory() as d:
            pfad = os.path.join(d, "unter", "z.csv")
            zeilen = [{"Ebene": "3. OG", "OKS": "A-G03-_01", "Raum_ID": 12345, "Status": "sicher",
                       "Freigabe": "J", "Stempel_Name": "Büro"},
                      {"Ebene": "3. OG", "OKS": "A-G03-_02", "Raum_ID": 12346, "Status": "unsicher",
                       "Freigabe": "N", "Stempel_Name": "Flur"},
                      {"Ebene": "3. OG", "OKS": "A-G03-_03", "Raum_ID": "", "Freigabe": "J"}]
            rs.schreibe_csv(pfad, rs.ZUORDNUNG_SPALTEN, zeilen)
            gelesen, mel = rs.lese_zuordnungsliste(pfad)
            self.assertEqual(len(gelesen), 2)
            self.assertEqual(len(mel), 1)                       # freigegeben ohne Raum_ID
            self.assertTrue(gelesen[0]["_freigabe"])
            self.assertFalse(gelesen[1]["_freigabe"])
            self.assertEqual(gelesen[0]["Raum_ID"], 12345)
            self.assertEqual(gelesen[0]["Stempel_Name"], "Büro")

    def test_freigabe_werte(self):
        for w in ("J", "ja", " x ", "1"):
            self.assertTrue(rs.ist_freigabe(w))
        for w in ("N", "", None, "nein"):
            self.assertFalse(rs.ist_freigabe(w))


class TestExcelFreundlich(unittest.TestCase):
    def test_excel_text(self):
        self.assertEqual(rs.excel_text("1.06"), '="1.06"')
        self.assertEqual(rs.excel_text("4.10 "), '="4.10"')
        self.assertEqual(rs.excel_text("0.1"), '="0.1"')
        self.assertEqual(rs.excel_text("n.v."), "n.v.")
        self.assertEqual(rs.excel_text("4.06b"), "4.06b")
        self.assertEqual(rs.excel_text("304a"), "304a")
        self.assertEqual(rs.excel_text(None), "")

    def test_zelle_text_und_dezimal(self):
        self.assertEqual(rs.zelle_text('="1.06"'), "1.06")
        self.assertEqual(rs.zelle_text("abc"), "abc")
        self.assertEqual(rs.fmt_dezimal(17.6, 2), "17,60")
        self.assertEqual(rs.fmt_dezimal(None), "")
        self.assertEqual(rs.parse_zahl(rs.fmt_dezimal(17.64)), 17.64)

    def test_csv_rundlauf_mit_excel_text(self):
        with tempfile.TemporaryDirectory() as d:
            pfad = os.path.join(d, "z.csv")
            zeilen = [{"OKS": "A-G01-_1", "Stempel_Nummer": rs.excel_text("1.06"), "Raum_ID": 5,
                       "Freigabe": "J", "Raum_Flaeche": rs.fmt_dezimal(12.5)}]
            rs.schreibe_csv(pfad, ["OKS", "Stempel_Nummer", "Raum_ID", "Freigabe", "Raum_Flaeche"], zeilen)
            with open(pfad, encoding="utf-8-sig") as f:
                roh = f.read()
            self.assertIn('"=""1.06"""', roh)           # korrekt zitiert -> Excel zeigt Text
            gelesen, _ = rs.lese_zuordnungsliste(pfad)
            self.assertEqual(gelesen[0]["Stempel_Nummer"], "1.06")
            self.assertEqual(gelesen[0]["Raum_Flaeche"], "12,50")

    def test_stempel_aus_csv_statt_aus_verfaelschter_liste(self):
        csv_stempel = rs.Stempel("A-G01-_1", nummer="1.06", name="Büro", x=0, y=0)
        zeile = {"OKS": "A-G01-_1", "Stempel_Nummer": "01. Jun", "Stempel_Name": "Büro"}
        s, quelle = rs.stempel_fuer_zeile(zeile, {"A-G01-_1": csv_stempel})
        self.assertEqual((s.nummer, quelle), ("1.06", "csv"))
        s, quelle = rs.stempel_fuer_zeile({"OKS": "A-G01-_9", "Stempel_Nummer": '="2.5"',
                                           "Stempel_Name": "Flur"}, {})
        self.assertEqual((s.nummer, s.name, quelle), ("2.5", "Flur", "liste"))


class TestXlsx(unittest.TestCase):
    SPALTEN = ["Ebene", "OKS", "Stempel_Nummer", "Stempel_Name", "Raum_ID", "Raum_Flaeche",
               "Abweichung_Prozent", "Status", "Freigabe", "Bemerkung"]
    ZEILEN = [
        {"Ebene": "EG- OK FFB", "OKS": "A-G00-_01", "Stempel_Nummer": "1.06", "Stempel_Name": "Büro & <Co>",
         "Raum_ID": 3258026, "Raum_Flaeche": 17.64, "Abweichung_Prozent": 1.2, "Status": "sicher",
         "Freigabe": "J", "Bemerkung": ""},
        {"Ebene": "EG- OK FFB", "OKS": "A-G00-_02", "Stempel_Nummer": "n.v.", "Stempel_Name": 'Flur "1"',
         "Raum_ID": 3258027, "Raum_Flaeche": 3.0, "Abweichung_Prozent": None, "Status": "unsicher",
         "Freigabe": "N", "Bemerkung": "zwei\nZeilen"},
        {"Ebene": "EG- OK FFB", "OKS": "A-G00-_03", "Stempel_Nummer": "4.06b", "Stempel_Name": "ä",
         "Raum_ID": 3258028, "Raum_Flaeche": 1.25, "Abweichung_Prozent": 6.5, "Status": "unsicher",
         "Freigabe": "N", "Bemerkung": "x"},
    ]

    def _schreibe(self, d):
        pfad = os.path.join(d, "unter", "z.xlsx")
        rs.schreibe_xlsx(pfad, "Zuordnung", self.SPALTEN, self.ZEILEN,
                         rs.ZUORDNUNG_ZAHLENFORMAT, "Freigabe", lambda z: z["Status"] == "unsicher")
        return pfad

    def test_rundlauf_eigener_leser(self):
        with tempfile.TemporaryDirectory() as d:
            tab = rs.lese_xlsx(self._schreibe(d))
        self.assertEqual(tab[0], self.SPALTEN)
        self.assertEqual(len(tab), 4)
        self.assertEqual(tab[1][2], "1.06")                       # bleibt Text, kein Datum
        self.assertEqual(tab[1][3], "Büro & <Co>")
        self.assertEqual(tab[1][4], "3258026")
        self.assertEqual(tab[1][5], "17.64")
        self.assertEqual(tab[2][3], 'Flur "1"')
        self.assertEqual(tab[2][6], "")                           # None -> leere Zelle
        self.assertEqual(tab[2][9], "zwei\nZeilen")
        self.assertEqual(tab[3][2], "4.06b")

    def test_zuordnungsliste_aus_xlsx(self):
        with tempfile.TemporaryDirectory() as d:
            zeilen, mel = rs.lese_zuordnungsliste(self._schreibe(d))
        self.assertEqual(mel, [])
        self.assertEqual([z["Raum_ID"] for z in zeilen], [3258026, 3258027, 3258028])
        self.assertEqual([z["_freigabe"] for z in zeilen], [True, False, False])
        self.assertEqual(zeilen[0]["Stempel_Nummer"], "1.06")

    def test_stempel_aus_xlsx_ordner(self):
        stempel_csv, _ = rs.lese_stempel_ordner(os.path.dirname(BEISPIEL_CSV))
        tab = rs.lese_tabelle(BEISPIEL_CSV)
        with tempfile.TemporaryDirectory() as d:
            kopf = tab[0]
            zeilen = [dict(zip(kopf, r)) for r in tab[1:]]
            rs.schreibe_xlsx(os.path.join(d, "Stempel.xlsx"), "Daten", kopf, zeilen)
            # Ausgabelisten und Sperrdateien dürfen nicht als Stempel gelesen werden
            rs.schreibe_xlsx(os.path.join(d, "Zuordnungsliste_1.xlsx"), "Z", ["OKS"], [{"OKS": "A"}])
            rs.schreibe_xlsx(os.path.join(d, "~$Stempel.xlsx"), "Z", ["OKS"], [{"OKS": "A"}])
            stempel, mel = rs.lese_stempel_ordner(d)
        self.assertEqual(mel, [])
        self.assertEqual(len(stempel), 23)
        self.assertEqual([s.oks for s in stempel], [s.oks for s in stempel_csv])
        self.assertAlmostEqual(stempel[0].x, stempel_csv[0].x)
        self.assertEqual(stempel[0].nummer, stempel_csv[0].nummer)

    def test_excel_gespeicherte_datei_mit_gemeinsamen_strings_und_luecken(self):
        """Simuliert eine von Excel gespeicherte Datei: shared strings, übersprungene Zellen."""
        import zipfile
        with tempfile.TemporaryDirectory() as d:
            pfad = os.path.join(d, "e.xlsx")
            ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            with zipfile.ZipFile(pfad, "w") as z:
                z.writestr("xl/workbook.xml",
                           '<workbook xmlns="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                           '<sheets><sheet name="Tabelle1" sheetId="1" r:id="rId3"/></sheets></workbook>' % ns)
                z.writestr("xl/_rels/workbook.xml.rels",
                           '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                           '<Relationship Id="rId3" Type="x" Target="/xl/worksheets/sheet7.xml"/></Relationships>')
                z.writestr("xl/sharedStrings.xml",
                           '<sst xmlns="%s"><si><t>OKS</t></si><si><t>Freigabe</t></si>'
                           '<si><r><t>A-G01-</t></r><r><t>_1</t></r></si><si><t>J</t></si></sst>' % ns)
                z.writestr("xl/worksheets/sheet7.xml",
                           '<worksheet xmlns="%s"><sheetData>'
                           '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="C1" t="s"><v>1</v></c></row>'
                           '<row r="3"><c r="A3" t="s"><v>2</v></c><c r="B3"><v>3258026</v></c>'
                           '<c r="C3" t="s"><v>3</v></c></row></sheetData></worksheet>' % ns)
            tab = rs.lese_xlsx(pfad)
        self.assertEqual(tab, [["OKS", "", "Freigabe"], [], ["A-G01-_1", "3258026", "J"]])

    def test_spaltenbuchstaben(self):
        self.assertEqual([rs._spaltenbuchstabe(i) for i in (0, 25, 26, 27, 701, 702)],
                         ["A", "Z", "AA", "AB", "ZZ", "AAA"])
        self.assertEqual([rs._spaltenindex(b) for b in ("A", "Z", "AA", "ZZ", "AAA")], [0, 25, 26, 701, 702])

    def test_ungueltige_xml_zeichen(self):
        with tempfile.TemporaryDirectory() as d:
            pfad = os.path.join(d, "x.xlsx")
            rs.schreibe_xlsx(pfad, "A/B", ["T"], [{"T": "a\x01b\x0bc"}])
            self.assertEqual(rs.lese_xlsx(pfad), [["T"], ["abc"]])

    def test_csv_zeilen_aufbereitung(self):
        z = rs.csv_zeilen([{"Stempel_Nummer": "1.06", "Raum_Flaeche": 12.5, "Abweichung_Prozent": None}],
                          rs.ZUORDNUNG_TEXTSPALTEN, rs.ZUORDNUNG_ZAHLENFORMAT)
        self.assertEqual(z[0], {"Stempel_Nummer": '="1.06"', "Raum_Flaeche": "12,50", "Abweichung_Prozent": None})

    def test_openpyxl_prueft_das_format(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl nicht installiert")
        with tempfile.TemporaryDirectory() as d:
            pfad = self._schreibe(d)
            wb = openpyxl.load_workbook(pfad)
            ws = wb.active
            self.assertEqual(ws.title, "Zuordnung")
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws.auto_filter.ref, "A1:J4")
            self.assertEqual(ws["C2"].value, "1.06")
            self.assertEqual(ws["C2"].data_type, "s")             # Text, kein Datum
            self.assertEqual(ws["E2"].value, 3258026)
            self.assertAlmostEqual(ws["F2"].value, 17.64)
            self.assertEqual(ws["F2"].number_format, "0.00")
            self.assertEqual(ws["G3"].value, None)
            self.assertTrue(ws["A1"].font.b)
            self.assertEqual(ws["A3"].fill.fgColor.rgb, "FFFFF2CC")   # unsicher markiert
            self.assertEqual(ws.data_validations.dataValidation[0].formula1, '"J,N"')
            wb.close()


class TestListenAusgabe(unittest.TestCase):
    def test_geprueftes_schreiben_und_ausweichen(self):
        meldungen = []

        class L(object):
            def __call__(self, t=""):
                meldungen.append(t)
        with tempfile.TemporaryDirectory() as d:
            erfolg = rs._schreibe_liste_geprueft(
                d, "z.csv", lambda pfad: rs.schreibe_csv(pfad, ["A", "B"], [{"A": 1, "B": "ü"}]), L())
            self.assertTrue(os.path.isfile(erfolg))
            self.assertIn("Bytes", meldungen[-1])
            # Ausgabeordner ist eine Datei -> Ausweichordner
            blockiert = os.path.join(d, "blockiert")
            with open(blockiert, "w"):
                pass
            erfolg2 = rs._schreibe_liste_geprueft(
                blockiert, "z2.csv", lambda pfad: rs.schreibe_csv(pfad, ["A"], [{"A": 1}]), L())
            self.assertTrue(erfolg2 is None or os.path.isfile(erfolg2))
            self.assertTrue(any("WARNUNG" in m for m in meldungen))
            if erfolg2:
                os.remove(erfolg2)


class TestGesamtbeispiel(unittest.TestCase):
    def test_beispiel_g03_synthetische_raeume(self):
        """Alle 23 Stempel der Beispieldatei; Räume mit den Stempelflächen (+1 %) ohne Treffer
        -> reine Flächenzuordnung muss alle 23 zuordnen, ohne Mehrfachvergabe."""
        stempel, _ = rs.lese_stempel_ordner(os.path.dirname(BEISPIEL_CSV))
        raeume = {1000 + i: {"flaeche": s.flaeche * 1.01, "pos": (s.x + 1, s.y + 1)}
                  for i, s in enumerate(stempel)}
        pos = {i: (s.x, s.y) for i, s in enumerate(stempel)}
        e = rs.ordne_zu(stempel, raeume, {i: [] for i in range(len(stempel))}, pos,
                        max_abstand=10)
        self.assertEqual(len(e["zuordnungen"]), 23)
        self.assertEqual(len({z.raum_id for z in e["zuordnungen"]}), 23)
        for z in e["zuordnungen"]:
            self.assertEqual(z.raum_id, 1000 + z.stempel_idx)   # jeder Stempel bekommt seinen Raum


if __name__ == "__main__":
    unittest.main()
